"""勤怠CSV自動集計ツール

入力CSV（1行＝1勤務ブロック）:
    社員名, 日付, 出勤, 退勤, 休憩(分), 現場名

出力Excel（3シート）:
    人別集計 / 日別明細 / 警告一覧

使い方:
    python attendance.py 入力.csv 出力.xlsx [--config config.json]

ルール（初期値。config.json で変更可）:
    - 総労働 = 退勤 - 出勤 - 休憩（同じ社員・同じ日の全ブロックを合算）
    - 退勤が出勤より前なら翌日扱い（日跨ぎ）。勤務は「出勤した日」に紐づける
    - 残業 = 1日の総労働が8時間を超えた分
    - 深夜帯時間 = 22:00〜翌5:00との重複時間。休憩の時間帯が入力にないため、休憩は控除しない
    - 時間帯の重複は、入力値をそのまま二重に合算した上で警告する
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from collections import defaultdict
from dataclasses import dataclass, fields
from datetime import date, datetime, timedelta
from pathlib import Path

# ---- 警告の種類 --------------------------------------------------------
W_DEFECT = "入力不備（空欄・形式不正）"
W_ZERO = "出勤＝退勤"
W_OVERLAP = "時間帯重複"
W_CROSSDAY = "日跨ぎ勤務"
W_LONG = "長時間勤務"

OVERLAP_NOTE = "⚠ 時間帯重複あり。総労働時間は入力値をそのまま合算しています。"


# ---- 設定 ---------------------------------------------------------------
@dataclass
class Config:
    night_start: str = "22:00"
    night_end: str = "05:00"
    daily_limit_hours: float = 8      # これを超えた分が残業
    long_day_hours: float = 12        # これを超えたら警告

    @property
    def daily_limit_min(self) -> int:
        return round(self.daily_limit_hours * 60)

    @property
    def long_day_min(self) -> int:
        return round(self.long_day_hours * 60)


class ConfigError(ValueError):
    """設定ファイルの内容が不正なとき。メッセージはそのまま利用者に表示できる。"""


_TIME_KEYS = ("night_start", "night_end")
_HOURS_KEYS = ("daily_limit_hours", "long_day_hours")


def validate_config(cfg: Config) -> None:
    """明らかな誤りだけを弾く。運用としてあり得る設定は禁止しない。

    - 時刻は HH:MM。night_start と night_end が同じ時刻なら深夜帯が成立しない
    - 時間数は 0 以上の有限の数値（0 は「全時間を残業扱い」、24超は「残業なし」の運用）
    - long_day_hours と daily_limit_hours は独立した閾値で、大小関係は問わない
    """
    parsed = {}
    for key in _TIME_KEYS:
        value = getattr(cfg, key)
        try:
            if not isinstance(value, str):
                raise ValueError
            parsed[key] = _parse_clock(value)
        except ValueError:
            raise ConfigError(f"{key}: 時刻は \"HH:MM\" 形式の文字列で指定してください（現在: {value!r}）") from None
    if parsed["night_start"] == parsed["night_end"]:
        raise ConfigError(f"night_start と night_end が同じ時刻です（{cfg.night_start}）。深夜帯の開始と終了は別の時刻にしてください")
    for key in _HOURS_KEYS:
        value = getattr(cfg, key)
        if (isinstance(value, bool) or not isinstance(value, (int, float))
                or not math.isfinite(value) or value < 0):
            raise ConfigError(f"{key}: 0以上の数値で指定してください（現在: {value!r}）")


def load_config(path: str | Path | None = None) -> Config:
    cfg = Config()
    if path:
        try:
            data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
        except json.JSONDecodeError as e:
            raise ConfigError(f"設定ファイルをJSONとして読めません（{e.lineno}行目）: {e.msg}") from e
        if not isinstance(data, dict):
            raise ConfigError("設定ファイルは { } で囲んだ形式（JSONオブジェクト）にしてください")
        known = [f.name for f in fields(Config)]
        unknown = [k for k in data if k not in known and not k.startswith("_")]
        if unknown:                      # 打ち間違いが黙って無視されるのを防ぐ
            raise ConfigError(f"未知の設定キー: {', '.join(unknown)}"
                              f"（使えるキー: {', '.join(known)}。先頭が _ のキーはメモ用として無視されます）")
        for key in known:
            if key in data:
                setattr(cfg, key, data[key])
    validate_config(cfg)
    return cfg


# ---- データ構造 ---------------------------------------------------------
@dataclass
class Block:
    row_no: int
    name: str
    day: date              # 出勤した日
    start: datetime
    end: datetime
    break_min: int
    site: str
    work_min: int          # 退勤-出勤-休憩（分）
    night_min: int         # 深夜帯との重複（休憩控除前・分）


@dataclass
class Daily:
    name: str
    day: date
    blocks: int = 0
    total_min: int = 0
    overtime_min: int = 0
    night_min: int = 0


@dataclass
class Warning:
    name: str
    day: date
    kind: str
    detail: str


@dataclass
class Result:
    daily: dict          # (name, date) -> Daily
    warnings: list       # list[Warning]

    def warning_kinds(self, name: str, day: date) -> set:
        return {w.kind for w in self.warnings if w.name == name and w.day == day}

    def person_summary(self) -> list:
        """人別集計: [(name, 出勤日数, 総労働, 残業, 深夜帯)] 分単位"""
        agg: dict = defaultdict(lambda: [0, 0, 0, 0])
        for (name, _), d in self.daily.items():
            a = agg[name]
            a[0] += 1
            a[1] += d.total_min
            a[2] += d.overtime_min
            a[3] += d.night_min
        return [(name, *vals) for name, vals in sorted(agg.items())]


# ---- 時刻計算 -----------------------------------------------------------
def fmt_hm(minutes: int) -> str:
    return f"{minutes // 60}:{minutes % 60:02d}"


def _parse_date(text: str) -> date:
    return datetime.strptime(text.strip().replace("/", "-"), "%Y-%m-%d").date()


def _parse_clock(text: str):
    return datetime.strptime(text.strip(), "%H:%M").time()


def night_minutes(start: datetime, end: datetime, cfg: Config) -> int:
    """[start, end] と深夜帯（毎日 night_start〜翌 night_end）の重複分。"""
    ns, ne = _parse_clock(cfg.night_start), _parse_clock(cfg.night_end)
    crosses = ne <= ns          # 22:00→05:00 のように日をまたぐ設定か
    total = timedelta(0)
    d = start.date() - timedelta(days=1)
    while d <= end.date():
        w_start = datetime.combine(d, ns)
        w_end = datetime.combine(d + timedelta(days=1 if crosses else 0), ne)
        lo, hi = max(start, w_start), min(end, w_end)
        if hi > lo:
            total += hi - lo
        d += timedelta(days=1)
    return int(total.total_seconds() // 60)


# ---- 読み込み -----------------------------------------------------------
def read_blocks(csv_path: str | Path, cfg: Config):
    """CSVを読み、(有効なBlockのリスト, 入力不備の警告リスト) を返す。"""
    blocks: list = []
    warnings: list = []
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader, start=2):      # 1行目はヘッダー
            name = (row.get("社員名") or "").strip()
            day_txt = (row.get("日付") or "").strip()
            in_txt = (row.get("出勤") or "").strip()
            out_txt = (row.get("退勤") or "").strip()
            brk_txt = (row.get("休憩(分)") or "").strip()
            site = (row.get("現場名") or "").strip()
            if not name and not day_txt and not in_txt and not out_txt:
                continue                                # 完全な空行は無視

            def defect(reason: str, day=None):
                warnings.append(Warning(
                    name or "(氏名なし)", day or date.min, W_DEFECT,
                    f"{i}行目: {reason}。この行は計算対象外です"))

            try:
                day = _parse_date(day_txt)
            except ValueError:
                defect("日付が空欄または形式不正")
                continue
            if not in_txt or not out_txt:
                defect("出勤または退勤が空欄", day)
                continue
            try:
                t_in, t_out = _parse_clock(in_txt), _parse_clock(out_txt)
                brk = int(brk_txt) if brk_txt else 0
            except ValueError:
                defect("出勤・退勤・休憩の形式不正", day)
                continue

            start = datetime.combine(day, t_in)
            end = datetime.combine(day, t_out)
            if end < start:                              # 日跨ぎ
                end += timedelta(days=1)
            work = int((end - start).total_seconds() // 60) - brk
            if brk < 0 or work < 0:
                defect("休憩が不正（勤務時間より長いなど）", day)
                continue
            blocks.append(Block(i, name, day, start, end, brk, site, work,
                                night_minutes(start, end, cfg)))
    return blocks, warnings


# ---- 集計・警告 ---------------------------------------------------------
def analyze(blocks: list, cfg: Config, input_warnings: list | None = None) -> Result:
    daily: dict = {}
    for b in blocks:
        d = daily.setdefault((b.name, b.day), Daily(b.name, b.day))
        d.blocks += 1
        d.total_min += b.work_min
        d.night_min += b.night_min
    for d in daily.values():
        d.overtime_min = max(0, d.total_min - cfg.daily_limit_min)

    warn: dict = defaultdict(list)       # (name, day, kind) -> [detail,...]

    for w in input_warnings or []:
        warn[(w.name, w.day, w.kind)].append(w.detail)

    for b in blocks:
        if b.end == b.start:
            warn[(b.name, b.day, W_ZERO)].append(
                f"{b.start:%H:%M}〜{b.end:%H:%M}（0時間として扱いました）")
        elif b.end.date() > b.start.date():
            warn[(b.name, b.day, W_CROSSDAY)].append(
                f"{b.start:%H:%M}〜翌{b.end:%H:%M}")

    for d in daily.values():
        if d.total_min > cfg.long_day_min:
            warn[(d.name, d.day, W_LONG)].append(
                f"総労働{fmt_hm(d.total_min)}（{cfg.long_day_hours:g}時間超）")

    # 重複チェック: 日付ではなく実際の日時で比べる（前日からの日跨ぎ勤務との重なりも拾う）
    by_person: dict = defaultdict(list)
    for b in blocks:
        if b.end > b.start:
            by_person[b.name].append(b)
    for name, lst in by_person.items():
        lst.sort(key=lambda b: (b.start, b.end))
        for i, a in enumerate(lst):
            for b in lst[i + 1:]:
                if b.start >= a.end:
                    break
                lo, hi = max(a.start, b.start), min(a.end, b.end)
                text = f"{lo:%m/%d %H:%M}〜{hi:%m/%d %H:%M} が重複"
                for day in {a.day, b.day}:
                    if text not in warn[(name, day, W_OVERLAP)]:
                        warn[(name, day, W_OVERLAP)].append(text)

    warnings = [Warning(n, d, k, " / ".join(details))
                for (n, d, k), details in warn.items()]
    warnings.sort(key=lambda w: (w.name, w.day, w.kind))
    return Result(daily, warnings)


def process(csv_path: str | Path, cfg: Config | None = None) -> Result:
    cfg = cfg or Config()
    blocks, input_warnings = read_blocks(csv_path, cfg)
    return analyze(blocks, cfg, input_warnings)


# ---- Excel出力 ----------------------------------------------------------
def write_excel(result: Result, out_path: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    TIME_FMT, DATE_FMT = "[h]:mm", "yyyy-mm-dd"
    NIGHT_HEAD = "深夜帯時間（休憩控除前）"

    wb = Workbook()

    def make_sheet(ws, headers, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = "A2"

    def style_rows(ws, time_cols=(), date_cols=(), start_row=2):
        for row in ws.iter_rows(min_row=start_row):
            for c in row:
                c.font = font
                if c.column in time_cols:
                    c.number_format = TIME_FMT
                if c.column in date_cols:
                    c.number_format = DATE_FMT

    # 1. 人別集計
    ws1 = wb.active
    ws1.title = "人別集計"
    make_sheet(ws1, ["社員名", "出勤日数", "総労働", "残業", NIGHT_HEAD], [14, 10, 12, 12, 26])
    for name, days, total, ot, night in result.person_summary():
        ws1.append([name, days, total / 1440, ot / 1440, night / 1440])
    style_rows(ws1, time_cols=(3, 4, 5))
    note_row = ws1.max_row + 2
    ws1.cell(note_row, 1, "※深夜帯時間は22:00〜翌5:00との重複時間で、休憩は控除していません"
                          "（休憩の時間帯が入力にないため）。給与計算に使う前にご確認ください。").font = font
    ws1.cell(note_row + 1, 1, "※総労働は入力値をそのまま合算しています（時間帯重複があっても補正しません）。").font = font

    # 2. 日別明細
    ws2 = wb.create_sheet("日別明細")
    make_sheet(ws2, ["社員名", "日付", "ブロック数", "総労働", "残業", NIGHT_HEAD, "備考"],
               [14, 12, 12, 12, 12, 26, 60])
    for (name, day), d in sorted(result.daily.items()):
        kinds = result.warning_kinds(name, day)
        notes = sorted(kinds - {W_OVERLAP})
        if W_OVERLAP in kinds:
            notes.append(OVERLAP_NOTE)
        ws2.append([name, day, d.blocks, d.total_min / 1440, d.overtime_min / 1440,
                    d.night_min / 1440, " / ".join(notes)])
        if kinds:
            for c in ws2[ws2.max_row]:
                c.fill = warn_fill
    style_rows(ws2, time_cols=(4, 5, 6), date_cols=(2,))

    # 3. 警告一覧
    ws3 = wb.create_sheet("警告一覧")
    make_sheet(ws3, ["社員名", "日付", "種別", "内容"], [14, 12, 24, 70])
    if result.warnings:
        for w in result.warnings:
            ws3.append([w.name, None if w.day == date.min else w.day, w.kind, w.detail])
    else:
        ws3.append(["警告はありません"])
    style_rows(ws3, date_cols=(2,))

    wb.save(out_path)


# ---- コマンドライン -----------------------------------------------------
def main(argv=None) -> int:
    p = argparse.ArgumentParser(description="勤怠CSV自動集計ツール")
    p.add_argument("input_csv")
    p.add_argument("output_xlsx")
    p.add_argument("--config", help="設定ファイル(config.json)")
    args = p.parse_args(argv)

    cfg = load_config(args.config)
    result = process(args.input_csv, cfg)
    write_excel(result, args.output_xlsx)
    print(f"集計完了: {len(result.daily)}件の日別明細、警告{len(result.warnings)}件 → {args.output_xlsx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
