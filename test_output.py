"""Excel出力とCLIの動作確認。

集計値そのものの検証は test_attendance.py。ここでは「出力ファイルが
仕様どおりの形で作られるか」だけを見る。openpyxl が無ければスキップ。

実行: python -m unittest -v
"""
import contextlib
import io
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

import attendance as at

try:
    from openpyxl import load_workbook
except ImportError:                                   # pragma: no cover
    load_workbook = None

ROOT = Path(__file__).parent
SAMPLE = ROOT / "sample_attendance.csv"
HEADER = "社員名,日付,出勤,退勤,休憩(分),現場名\n"


@unittest.skipIf(load_workbook is None, "openpyxl が未インストール")
class TestExcelOutput(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)

    def export(self, csv_path: Path):
        out = self.tmp / "out.xlsx"
        at.write_excel(at.process(csv_path, at.Config()), out)
        return load_workbook(out)

    def test_sheets_and_headers(self):
        wb = self.export(SAMPLE)
        self.assertEqual(wb.sheetnames, ["人別集計", "日別明細", "警告一覧"])
        self.assertEqual([c.value for c in wb["人別集計"][1]],
                         ["社員名", "出勤日数", "総労働", "残業", "深夜帯時間（休憩控除前）"])
        self.assertEqual([c.value for c in wb["日別明細"][1]],
                         ["社員名", "日付", "ブロック数", "総労働", "残業", "深夜帯時間（休憩控除前）", "備考"])
        self.assertEqual([c.value for c in wb["警告一覧"][1]], ["社員名", "日付", "種別", "内容"])

    def test_durations_are_written_as_excel_times(self):
        # 山田 10/2 は総労働9:30・残業1:30（test_attendance.py の正解と同じ）
        ws = self.export(SAMPLE)["日別明細"]
        row = next(r for r in ws.iter_rows(min_row=2)
                   if r[0].value == "山田" and r[1].value == datetime(2026, 10, 2))
        # [h]:mm 書式のセルは openpyxl が timedelta として返す
        self.assertEqual(row[3].value, timedelta(hours=9, minutes=30))
        self.assertEqual(row[4].value, timedelta(minutes=90))
        self.assertEqual(row[3].number_format, "[h]:mm")

    def test_rows_with_warnings_are_highlighted(self):
        ws = self.export(SAMPLE)["日別明細"]
        by_name = {r[0].value: r for r in ws.iter_rows(min_row=2)}
        self.assertTrue(by_name["高橋"][0].fill.fgColor.rgb.endswith("FFF2CC"))   # 時間帯重複あり
        self.assertFalse(by_name["山田"][0].fill.fgColor.rgb.endswith("FFF2CC"))  # 警告なし

    def test_no_warnings_message(self):
        csv_path = self.tmp / "clean.csv"
        csv_path.write_text(HEADER + "山田,2026-10-01,09:00,17:00,0,現場A\n", encoding="utf-8")
        ws = self.export(csv_path)["警告一覧"]
        self.assertEqual(ws["A2"].value, "警告はありません")

    def test_cli_main_writes_file(self):
        out = self.tmp / "cli.xlsx"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            code = at.main([str(SAMPLE), str(out), "--config", str(ROOT / "config.json")])
        self.assertEqual(code, 0)
        self.assertTrue(out.exists())
        self.assertIn("集計完了", buf.getvalue())

    def test_cli_rejects_bad_config(self):
        bad = self.tmp / "bad.json"
        bad.write_text('{"daily_limit_hour": 7}', encoding="utf-8")
        with self.assertRaises(at.ConfigError):
            at.main([str(SAMPLE), str(self.tmp / "x.xlsx"), "--config", str(bad)])
        self.assertFalse((self.tmp / "x.xlsx").exists())    # 失敗時に出力を作らない


if __name__ == "__main__":
    unittest.main()
